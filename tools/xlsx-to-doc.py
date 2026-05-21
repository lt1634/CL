from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import openpyxl

# Load workbook
wb = openpyxl.load_workbook('/Users/timnewmac/Library/CloudStorage/GoogleDrive-bfhywt@bfhmc.edu.hk/我的雲端硬碟/Documents/創業/2026_一人公司創業關鍵詞手冊.xlsx')

doc = Document()

# ── Title ──────────────────────────────────────────────
title = doc.add_heading('一人公司創業關鍵詞手冊', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph('2026 最新版｜100 個必學關鍵詞').alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph()

# ── Style helpers ──────────────────────────────────────
def add_term(doc, number, en, zh, category, explanation):
    # Heading: number + English term
    h = doc.add_heading(f'{number}. {en}', level=2)
    # Chinese name + category in accent
    p = doc.add_paragraph()
    run = p.add_run(f'{zh}  |  分類：{category}')
    run.italic = True
    run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
    run.font.size = Pt(10)
    # Explanation
    p2 = doc.add_paragraph(explanation)
    p2.style = 'Quote'
    doc.add_paragraph()

# ── Sheet 1: 發展趨勢分析 ──────────────────────────────
doc.add_heading('2026 發展趨勢分析', level=1)

trend_sheet = wb['2026 發展趨勢分析']
for row in trend_sheet.iter_rows(min_row=2, values_only=True):
    if row[0] and row[1]:
        label = str(row[0]).strip()
        content = str(row[1]).strip()
        p = doc.add_paragraph()
        run = p.add_run(f'{label}：')
        run.bold = True
        p.add_run(content)
        doc.add_paragraph()

doc.add_page_break()

# ── Sheet 2: 100個必學關鍵詞 ─────────────────────────
doc.add_heading('100 個必學關鍵詞', level=1)

keywords_sheet = wb['100個必學關鍵詞']
for row in keywords_sheet.iter_rows(min_row=2, values_only=True):
    if row[0] is None:
        continue
    number = row[0]
    en     = row[1] or ''
    zh     = row[2] or ''
    cat    = row[3] or ''
    expl   = row[4] or ''
    add_term(doc, number, en, zh, cat, expl)

# ── Save ──────────────────────────────────────────────
out = '/Users/timnewmac/Desktop/CL/tools/一人公司創業關鍵詞手冊.docx'
doc.save(out)
print(f'Saved → {out}')
